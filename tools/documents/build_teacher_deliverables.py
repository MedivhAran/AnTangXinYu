from __future__ import annotations

import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
SOURCE_DIR = ROOT / "docs" / "deliverables" / "teacher-review-2026-08"
ASSET_DIR = SOURCE_DIR / "assets"
OUTPUT_DIR = ROOT / "deliverables" / "teacher-review-2026-08"

PAGE_WIDTH_DXA = 12240
PAGE_HEIGHT_DXA = 15840
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120

LATIN_FONT = "Calibri"
CJK_FONT = "Microsoft YaHei"
MONO_FONT = "Consolas"

BRAND = "11866F"
BRAND_BRIGHT = "20B894"
BRAND_PALE = "E7F7F2"
INK = "17302E"
MUTED = "667875"
BORDER = "D9E5E2"
LIGHT_GRAY = "F2F4F7"
CAUTION = "FFF6DD"
RISK_PALE = "FCEDEA"
RISK = "B34840"


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def set_run_font(
    run,
    *,
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
    latin: str = LATIN_FONT,
    east_asia: str = CJK_FONT,
) -> None:
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_style_font(style, *, size: float, color: str, bold: bool = False) -> None:
    style.font.name = LATIN_FONT
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), LATIN_FONT)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), LATIN_FONT)
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), CJK_FONT)
    style.font.size = Pt(size)
    style.font.color.rgb = rgb(color)
    style.font.bold = bold


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, *, top: int = 80, bottom: int = 80, start: int = 120, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_table_geometry(table, widths: list[int]) -> None:
    if sum(widths) != CONTENT_WIDTH_DXA:
        raise ValueError(f"table widths must total {CONTENT_WIDTH_DXA}: {widths}")
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")


def add_bottom_border(paragraph, color: str = BRAND_BRIGHT, size: int = 12) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_left_border(paragraph, color: str = BRAND_BRIGHT, size: int = 18) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), color)
    p_bdr.append(left)


def set_paragraph_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run_node = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), LATIN_FONT)
    fonts.set(qn("w:hAnsi"), LATIN_FONT)
    fonts.set(qn("w:eastAsia"), CJK_FONT)
    props.append(fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), MUTED)
    props.append(color)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "18")
    props.append(size)
    run_node.append(props)
    text = OxmlElement("w:t")
    text.text = "1"
    run_node.append(text)
    field.append(run_node)
    paragraph._p.append(field)
    run = paragraph.add_run(" 页")
    set_run_font(run, size=9, color=MUTED)


def add_numbering(doc: Document, *, num_id: int, abstract_id: int, fmt: str, marker: str) -> None:
    numbering = doc.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), fmt)
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), marker)
    lvl.append(lvl_text)
    suffix = OxmlElement("w:suff")
    suffix.set(qn("w:val"), "tab")
    lvl.append(suffix)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    p_pr.append(tabs)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "720")
    indent.set(qn("w:hanging"), "360")
    p_pr.append(indent)
    lvl.append(p_pr)
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), LATIN_FONT)
    fonts.set(qn("w:hAnsi"), LATIN_FONT)
    fonts.set(qn("w:eastAsia"), CJK_FONT)
    r_pr.append(fonts)
    lvl.append(r_pr)
    abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)


def apply_num(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)
    p_pr.append(num_pr)


def configure_document(doc: Document, short_title: str) -> None:
    doc.settings.odd_and_even_pages_header_footer = True
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    set_style_font(normal, size=11, color=INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    h1 = doc.styles["Heading 1"]
    set_style_font(h1, size=16, color=BRAND, bold=True)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(8)
    h1.paragraph_format.keep_with_next = True

    h2 = doc.styles["Heading 2"]
    set_style_font(h2, size=13, color=BRAND, bold=True)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    h2.paragraph_format.keep_with_next = True

    h3 = doc.styles["Heading 3"]
    set_style_font(h3, size=12, color=INK, bold=True)
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(4)
    h3.paragraph_format.keep_with_next = True

    for header in (section.header, section.even_page_header):
        p = header.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
        left = p.add_run("安糖心语")
        set_run_font(left, size=9, color=MUTED, bold=True)
        right = p.add_run(f"\t{short_title}")
        set_run_font(right, size=9, color=MUTED)

    for footer in (section.footer, section.even_page_footer):
        footer_p = footer.paragraphs[0]
        footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer_p.paragraph_format.space_before = Pt(0)
        footer_p.paragraph_format.space_after = Pt(0)
        add_page_field(footer_p)

    add_numbering(doc, num_id=90, abstract_id=90, fmt="bullet", marker="•")
    add_numbering(doc, num_id=91, abstract_id=91, fmt="decimal", marker="%1.")


def add_inline(paragraph, text: str, *, size: float = 11, color: str = INK) -> None:
    token_pattern = re.compile(r"(\*\*.+?\*\*|`.+?`)")
    position = 0
    for match in token_pattern.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position : match.start()])
            set_run_font(run, size=size, color=color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, color=color, bold=True)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(
                run,
                size=max(9, size - 0.5),
                color=BRAND,
                latin=MONO_FONT,
                east_asia=CJK_FONT,
            )
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run_font(run, size=size, color=color)


def add_body_paragraph(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10
    add_inline(p, text)


def add_list_item(doc: Document, text: str, *, ordered: bool) -> None:
    p = doc.add_paragraph()
    apply_num(p, 91 if ordered else 90)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    p.paragraph_format.keep_together = True
    add_inline(p, text)


def add_callout(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.10
    add_left_border(p)
    set_paragraph_shading(p, BRAND_PALE if not text.startswith("[") else LIGHT_GRAY)
    add_inline(p, text, size=10.5, color=INK if not text.startswith("[") else MUTED)


def table_widths(headers: list[str]) -> list[int]:
    count = len(headers)
    if count == 2:
        return [3000, 6360]
    if count == 3 and any("状态" in header for header in headers):
        return [2700, 1500, 5160]
    if count == 3:
        return [2200, 3280, 3880]
    if count == 4:
        return [1700, 2500, 2480, 2680]
    base = CONTENT_WIDTH_DXA // count
    widths = [base] * count
    widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
    return widths


def clean_table_cell(text: str) -> str:
    return text.strip().replace("**", "").replace("`", "")


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    columns = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"
    set_table_geometry(table, table_widths(rows[0]))
    set_repeat_table_header(table.rows[0])

    for row_index, source_row in enumerate(rows):
        for column_index, value in enumerate(source_row):
            cell = table.cell(row_index, column_index)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.0
            if column_index == 1 and len(rows[0]) == 3:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(clean_table_cell(value))
            set_run_font(
                run,
                size=9.3,
                color=INK,
                bold=row_index == 0,
            )
            if row_index == 0:
                set_cell_shading(cell, BRAND_PALE)
            elif column_index == 1:
                status_text = clean_table_cell(value)
                if status_text == "未实现":
                    set_cell_shading(cell, RISK_PALE)
                    run.font.color.rgb = rgb(RISK)
                elif status_text == "当前版本不采用":
                    set_cell_shading(cell, LIGHT_GRAY)
                    run.font.color.rgb = rgb(MUTED)
                elif "部分" in status_text:
                    set_cell_shading(cell, CAUTION)
                elif "完成" in status_text:
                    set_cell_shading(cell, BRAND_PALE)
                    run.font.color.rgb = rgb(BRAND)

    after = doc.add_paragraph()
    after.paragraph_format.space_before = Pt(0)
    after.paragraph_format.space_after = Pt(2)


def add_code_block(doc: Document, lines: list[str]) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.18)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.0
    set_paragraph_shading(p, LIGHT_GRAY)
    run = p.add_run("\n".join(lines))
    set_run_font(run, size=9, color=INK, latin=MONO_FONT, east_asia=CJK_FONT)


def add_figure(doc: Document, source_path: Path, alt: str, figure_number: int) -> None:
    if not source_path.exists():
        add_callout(doc, f"[待生成图片：{alt}]")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run()
    inline = run.add_picture(str(source_path), width=Inches(6.15))
    inline._inline.docPr.set("descr", alt)
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(0)
    caption.paragraph_format.space_after = Pt(10)
    cap_run = caption.add_run(f"图 {figure_number}  {alt}")
    set_run_font(cap_run, size=9, color=MUTED)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    raw_rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
        raw_rows.append(cells)
        index += 1
    rows = [
        row
        for row in raw_rows
        if not all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in row)
    ]
    return rows, index


def read_markdown(path: Path) -> tuple[str, list[str], list[str]]:
    lines = path.read_text(encoding="utf-8").splitlines()
    title = next(line[2:].strip() for line in lines if line.startswith("# "))
    first_section = next(index for index, line in enumerate(lines) if line.startswith("## "))
    metadata = [
        line.rstrip("  ").strip()
        for line in lines[1:first_section]
        if line.strip()
    ]
    return title, metadata, lines[first_section:]


def add_masthead(doc: Document, title: str, metadata: list[str]) -> None:
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(14)
    kicker.paragraph_format.space_after = Pt(4)
    run = kicker.add_run("安糖心语 · 阶段交付")
    set_run_font(run, size=10, color=BRAND_BRIGHT, bold=True)

    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(8)
    title_p.paragraph_format.keep_with_next = True
    run = title_p.add_run(title)
    set_run_font(run, size=24, color=INK, bold=True)

    for item in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        if "：" in item:
            label, value = item.split("：", 1)
            label_run = p.add_run(f"{label}：")
            set_run_font(label_run, size=10.5, color=MUTED, bold=True)
            value_run = p.add_run(value)
            set_run_font(value_run, size=10.5, color=MUTED)
        else:
            add_inline(p, item, size=10.5, color=MUTED)

    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(6)
    rule.paragraph_format.space_after = Pt(8)
    add_bottom_border(rule)


def render_markdown(source: Path, output: Path, short_title: str) -> None:
    title, metadata, lines = read_markdown(source)
    doc = Document()
    configure_document(doc, short_title)
    add_masthead(doc, title, metadata)
    doc.core_properties.title = title
    doc.core_properties.author = "张旭睿"
    doc.core_properties.subject = "安糖心语项目阶段材料"

    index = 0
    figure_number = 1
    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()
        if not stripped:
            index += 1
            continue
        if stripped.startswith("```"):
            code: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code.append(lines[index])
                index += 1
            add_code_block(doc, code)
            index += 1
            continue
        image_match = re.fullmatch(r"!\[(.+?)\]\((.+?)\)", stripped)
        if image_match:
            source_path = (source.parent / image_match.group(2)).resolve()
            add_figure(doc, source_path, image_match.group(1), figure_number)
            figure_number += 1
            index += 1
            continue
        if stripped.startswith("|"):
            rows, index = parse_table(lines, index)
            add_table(doc, rows)
            continue
        heading_match = re.match(r"^(#{2,4})\s+(.+)$", stripped)
        if heading_match:
            level = len(heading_match.group(1)) - 1
            p = doc.add_paragraph(style=f"Heading {level}")
            add_inline(p, heading_match.group(2), size={1: 16, 2: 13, 3: 12}[level], color=BRAND if level < 3 else INK)
            index += 1
            continue
        if stripped.startswith(">"):
            quote_lines: list[str] = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                quote_lines.append(lines[index].strip()[1:].strip())
                index += 1
            add_callout(doc, " ".join(quote_lines))
            continue
        if re.match(r"^\d+\.\s+", stripped):
            text = re.sub(r"^\d+\.\s+", "", stripped)
            add_list_item(doc, text, ordered=True)
            index += 1
            continue
        if stripped.startswith("- "):
            add_list_item(doc, stripped[2:].strip(), ordered=False)
            index += 1
            continue

        paragraph_lines = [stripped.rstrip("  ")]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if not candidate:
                break
            if (
                candidate.startswith("#")
                or candidate.startswith("|")
                or candidate.startswith(">")
                or candidate.startswith("- ")
                or candidate.startswith("```")
                or candidate.startswith("![")
                or re.match(r"^\d+\.\s+", candidate)
            ):
                break
            paragraph_lines.append(candidate.rstrip("  "))
            index += 1
        add_body_paragraph(doc, " ".join(paragraph_lines))

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def load_font(size: int, *, bold: bool = False):
    candidates = [
        Path(r"C:\Windows\Fonts\msyhbd.ttc" if bold else r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\simhei.ttf"),
        Path(r"C:\Windows\Fonts\arial.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def draw_centered_text(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], text: str, font, fill: str) -> None:
    left, top, right, bottom = box
    bbox = draw.multiline_textbbox((0, 0), text, font=font, spacing=10, align="center")
    width = bbox[2] - bbox[0]
    height = bbox[3] - bbox[1]
    draw.multiline_text(
        ((left + right - width) / 2, (top + bottom - height) / 2),
        text,
        font=font,
        fill=fill,
        spacing=10,
        align="center",
    )


def draw_box(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    title: str,
    detail: str,
    *,
    fill: str,
    outline: str = "#20B894",
) -> None:
    draw.rounded_rectangle(box, radius=28, fill=fill, outline=outline, width=4)
    left, top, right, _ = box
    title_font = load_font(29, bold=True)
    body_font = load_font(22)
    title_bbox = draw.textbbox((0, 0), title, font=title_font)
    title_width = title_bbox[2] - title_bbox[0]
    draw.text(((left + right - title_width) / 2, top + 28), title, font=title_font, fill="#17302E")
    draw_centered_text(draw, (left + 20, top + 78, right - 20, box[3] - 16), detail, body_font, "#49615E")


def arrow(
    draw: ImageDraw.ImageDraw,
    start: tuple[int, int],
    end: tuple[int, int],
    *,
    color: str = "#20B894",
    width: int = 6,
) -> None:
    draw.line((start, end), fill=color, width=width)
    x1, y1 = start
    x2, y2 = end
    dx = x2 - x1
    dy = y2 - y1
    length = max((dx * dx + dy * dy) ** 0.5, 1)
    ux, uy = dx / length, dy / length
    px, py = -uy, ux
    head = 18
    wing = 10
    points = [
        (x2, y2),
        (x2 - ux * head + px * wing, y2 - uy * head + py * wing),
        (x2 - ux * head - px * wing, y2 - uy * head - py * wing),
    ]
    draw.polygon(points, fill=color)


def build_architecture_diagram() -> None:
    canvas = Image.new("RGB", (1800, 1120), "#F5F9F8")
    draw = ImageDraw.Draw(canvas)
    title_font = load_font(42, bold=True)
    subtitle_font = load_font(22)
    draw.text((80, 48), "安糖心语当前运行架构", font=title_font, fill="#17302E")
    draw.text((80, 108), "以普通程序控制数据、规则与发送，Agent 只负责理解和候选内容", font=subtitle_font, fill="#667875")

    app_box = (70, 250, 405, 555)
    api_box = (510, 215, 955, 600)
    external_box = (1060, 145, 1710, 325)
    worker_box = (1060, 390, 1710, 630)
    db_box = (510, 780, 955, 1005)
    memory_box = (1060, 760, 1710, 1025)
    wearable_box = (70, 770, 405, 1010)

    draw_box(draw, app_box, "Android App", "连续聊天 · 健康档案\nHealth Connect · 系统通知", fill="#FFFFFF")
    draw_box(draw, api_box, "FastAPI", "身份与数据控制\nCore Agent 与工具编排\n流式回复 · 校验 · 审计", fill="#FFFFFF")
    draw_box(draw, external_box, "外部能力", "DeepSeek 大模型\nTavily 网页搜索与读取", fill="#F0FAF7")
    draw_box(draw, worker_box, "主动关怀 Worker", "持久任务 → 关怀 Agent\n→ Auditor → Expo/FCM", fill="#FFFFFF")
    draw_box(draw, db_box, "主 PostgreSQL", "账号 · 消息 · 健康档案\n手环观测 · 任务 · 推送记录", fill="#F0FAF7")
    draw_box(draw, memory_box, "Hindsight", "长期记忆 API\n专用 PostgreSQL 与检索索引", fill="#FFFFFF")
    draw_box(draw, wearable_box, "手环数据入口", "Active 2 → Zepp\n→ Health Connect → App", fill="#FFFFFF")

    arrow(draw, (405, 400), (510, 400))
    arrow(draw, (510, 455), (405, 455), color="#8AA8A2", width=4)
    arrow(draw, (955, 300), (1060, 245))
    arrow(draw, (735, 600), (735, 780))
    arrow(draw, (955, 520), (1060, 505))
    arrow(draw, (1280, 630), (925, 820), color="#8AA8A2", width=4)
    arrow(draw, (955, 875), (1060, 875))
    arrow(draw, (405, 880), (510, 520))
    arrow(draw, (1060, 940), (955, 550), color="#8AA8A2", width=4)

    label_font = load_font(18)
    draw.text((425, 365), "API / NDJSON", font=label_font, fill="#667875")
    draw.text((760, 675), "业务事实", font=label_font, fill="#667875")
    draw.text((970, 455), "领取任务", font=label_font, fill="#667875")
    draw.text((965, 838), "相关记忆", font=label_font, fill="#667875")

    note_font = load_font(22, bold=True)
    note = "关键边界：健康档案和设备记录由普通程序写入；Agent、Hindsight、规则和推送各自职责分离。"
    draw.text((80, 1062), note, font=note_font, fill="#11866F")

    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    canvas.save(ASSET_DIR / "current-system-architecture.png", format="PNG", optimize=True)


def main() -> None:
    build_architecture_diagram()
    jobs = [
        (
            "01_项目阶段进展与功能对照.md",
            "安糖心语_项目阶段进展与功能对照_V1.0_待补截图.docx",
            "项目阶段进展与功能对照",
        ),
        (
            "02_产品需求说明.md",
            "安糖心语_产品需求说明_V1.0.docx",
            "产品需求说明",
        ),
        (
            "03_系统设计说明.md",
            "安糖心语_系统设计说明_V1.0.docx",
            "系统设计说明",
        ),
    ]
    for source_name, output_name, short_title in jobs:
        render_markdown(
            SOURCE_DIR / source_name,
            OUTPUT_DIR / output_name,
            short_title,
        )


if __name__ == "__main__":
    main()
